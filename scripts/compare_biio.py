"""Independent cell-to-runtime comparison. templates.json is emitted by JS tests."""
import json,sys,pathlib,re
import openpyxl
from import_biio import text,CATS
root=pathlib.Path(sys.argv[1]);templates=json.loads(pathlib.Path(sys.argv[2]).read_text(encoding='utf-8'))
books={};differences=[];cells=plans=0
for t in templates:
    source=next(p for p in root.glob('*Mesociclo*.xlsx') if int(p.name[0])==CATS.index(t['categoria'])+1 and ('(GIRL)' in p.name)==(t['sexo']=='M'))
    wb=openpyxl.load_workbook(source,data_only=True);books[source.name]=True
    expected_sheets=[ws.title for ws in wb if ws.title.lower().startswith('entreno')]
    assert len([en for en in t['entrenos'] if not en['letra'].startswith('MX')])==len(expected_sheets)
    for en in t['entrenos']:
        if not en['letra'].startswith('MX'):
            ws=wb[en['ejercicios'][0]['source']['sheet']]
            header=next(r for r in range(1,25) if any(text(c)=='1º' for c in ws[r]))
            cols=[c.column for c in ws[header] if re.fullmatch(r'\d+º',text(c))]
            target=next(r for r in range(header+1,header+5) if 'TARGET' in text(ws.cell(r,1)))
            primary=[r for r in range(target+1,ws.max_row+1) if text(ws.cell(r,1)) and any(re.match(r'^(?:\d+(?:\s*o\s*\d+)?|\?)\s*[xX]',text(ws.cell(r,c))) for c in cols)]
            if primary!=[e['source']['row'] for e in en['ejercicios']]:differences.append([source.name,ws.title,'exercise coverage/order',primary])
            if len(cols)!=t['numMicro']:differences.append([source.name,'microcycle count'])
            for ej in en['ejercicios']:
                if set(ej['planByMicro'])!=set(map(str,range(1,len(cols)+1))):differences.append([source.name,ej['nombre'],'missing microcycle'])
            for mn,p in en.get('cardioByMicro',{}).items():
                expected=' / '.join(text(ws[c]) for c in p['sourceCells'])
                cells+=len(p['sourceCells'])
                if expected!=p['label']:differences.append([source.name,ws.title,mn,'cardio'])
        for ej in en['ejercicios']:
            ws=wb[ej['source']['sheet']];row=ej['source']['row']
            # Name/order comes from the independently loaded source, not the generated JSON.
            first=text(ws.cell(row,7 if en['letra'].startswith('MX') and text(ws.cell(row,7)) and not text(ws.cell(row,7)).startswith('Nota:') else 1))
            if not ej['nombre'].startswith(first):differences.append([t['sexo'],t['categoria'],en['letra'],ej['nombre'],'name',first])
            for mn,p in ej['planByMicro'].items():
                plans+=1
                for coord,actual in p.get('sourceCells',{}).items():
                    expected=text(ws[coord]);cells+=1
                    if actual!=expected:differences.append([source.name,ws.title,coord,expected,actual])
                rest=' / '.join(text(ws[c]) for c in p.get('sourceRestCells',{}));cells+=len(p.get('sourceRestCells',{}))
                if rest!=p['pausa']:differences.append([source.name,ws.title,row,mn,'rest',rest,p['pausa']])
                percent=' / '.join(text(ws[c]) for c in p.get('sourceCells',{}) if '%' in text(ws[c]))
                if percent!=p['porcentaje']:differences.append([source.name,ws.title,row,mn,'percentage',percent,p['porcentaje']])
                if p.get('sourceCells'):
                    expected=' / '.join(text(ws[c]) for c in p['sourceCells'])
                    if p['label']!=expected:differences.append([source.name,ws.title,row,mn,'label'])
                    # Check explicit set count independently; optional range remains literal.
                    first_value=next(iter(p['sourceCells'].values()))
                    nums=re.match(r'(\d+)(?:\s*o\s*(\d+))?\s*[xX]',first_value)
                    if nums and p['series']!=int(nums[2] or nums[1]):differences.append([source.name,ws.title,row,mn,'series'])
                    c=cols[int(mn)-1]
                    kg=next(r for r in range(row+1,ws.max_row+1) if text(ws.cell(r,c)).lower()=='kg')
                    prescribed={ws.cell(r,cc).coordinate:text(ws.cell(r,cc)) for r in range(row,kg) for cc in [c,c+1] if text(ws.cell(r,cc))}
                    if prescribed!=p['sourceCells']:differences.append([source.name,ws.title,row,mn,'missing prescription cells'])
                    inline=re.match(r'\d+(?:\s*o\s*\d+)?\s*x\s*(.*)',text(ws.cell(row,c)),re.I)
                    reps=([inline[1]] if inline and inline[1] else [])+[text(ws.cell(r,c+1)) for r in range(row,kg) if text(ws.cell(r,c+1)) and '%' not in text(ws.cell(r,c+1))]
                    if reps!=p['repsTarget']:differences.append([source.name,ws.title,row,mn,'repetitions',reps,p['repsTarget']])
report={'files':len(books),'templates':len(templates),'plans':plans,'cellsCompared':cells,'differences':differences,'scope':'Names, workbook variants, literal plan cells, explicit series and runtime labels. Blank prescriptions and variable counts remain literal.'}
print(json.dumps(report,ensure_ascii=False,indent=2));sys.exit(bool(differences))
