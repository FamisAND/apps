"""Read original BIIO cells, including merged-cell anchors. Never infer missing prescriptions."""
import json, re, hashlib, pathlib, argparse
import openpyxl

CATS = ['Reacondicionamiento','Preparación fuerza','Especialización técnica','Fuerza 1','Fuerza 2','Hibrido','Hipertrofia','Calidad muscular']

def text(cell):
    v = cell.value
    if v is None: return ''
    if isinstance(v, (int,float)) and '%' in cell.number_format:
        return f'{v*100:g}%'
    return str(v).strip()

def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    data = {'numMicro':0,'entrenos':[], 'sourceFile':path.name, 'sourceSha256':hashlib.sha256(path.read_bytes()).hexdigest()}
    diary=next((s for s in wb if s.title.lower()=='diario'),None)
    data['descripcion']='\n'.join(text(diary[c]) for c in ['A6','A9','N9','A12','F12'] if text(diary[c])) if diary else ''
    for ws in wb:
        if not ws.title.lower().startswith('entreno'): continue
        header = next(r for r in range(1,25) if any(text(c)=='1º' for c in ws[r]))
        cols = [c.column for c in ws[header] if re.fullmatch(r'\d+º',text(c))]
        num = len(cols)
        data['numMicro'] = max(data['numMicro'],num)
        titleRow = next(r for r in range(header+1,header+5) if 'TARGET' in text(ws.cell(r,1)))
        starts=[]
        for r in range(titleRow+1,ws.max_row+1):
            a=text(ws.cell(r,1))
            if not a or re.search(r'^(?:tiempo|pausa|\d+[ªº]|.*aer[oó]b|nota|\[|altern|en )',a,re.I): continue
            if any(re.match(r'^(?:\d+(?:\s*o\s*\d+)?|\?)\s*[xX]',text(ws.cell(r,c))) for c in cols): starts.append(r)
        assert starts, (path.name,ws.title)
        aerobic=next((r for r in range(starts[-1]+1,ws.max_row+1) if re.search(r'aer[oó]b',text(ws.cell(r,1)),re.I)),ws.max_row+1)
        en={'letra':ws.title[-1].upper(),'nombre':text(ws['A1']), 'ejercicios':[], 'indicaciones':{},'cardioByMicro':{}}
        for mn,c in enumerate(cols,1):
            en['indicaciones'][str(mn)]='\n'.join(text(ws.cell(r,cc)) for r in range(header+1,titleRow) for cc in [1,c] if text(ws.cell(r,cc)))
            cardioCells={ws.cell(r,cc).coordinate:text(ws.cell(r,cc)) for r in range(aerobic+1,ws.max_row+1) for cc in [1,c,c+1] if text(ws.cell(r,cc))}
            en['cardioByMicro'][str(mn)]={'label':' / '.join(cardioCells.values()),'sourceCells':cardioCells}
        for index,start in enumerate(starts):
            end=starts[index+1] if index+1<len(starts) else aerobic
            kg=next(r for r in range(start+1,end) if text(ws.cell(r,cols[0])).lower()=='kg')
            # Alternate exercises are explicitly introduced by an alternation/superset label.
            names=[text(ws.cell(start,1))]
            nameRows=[start]
            for r in range(start+1,kg):
                if re.search(r'altern|en (?:TRISERIE|SUPERSERIE)',text(ws.cell(r,1)),re.I):
                    if text(ws.cell(r+1,1)):
                        names.append(text(ws.cell(r+1,1))); nameRows.append(r+1)
            notes=[]
            for r in range(start+1,kg):
                for cc in range(1,cols[0]):
                    v=text(ws.cell(r,cc))
                    if v and r not in nameRows and not re.fullmatch(r'(?:Nota[s]?|Note)?[\s.…]*',v,re.I): notes.append(v)
            notes.extend(text(ws.cell(r,1)) for r in range(kg+1,end) if text(ws.cell(r,1)).startswith('+'))
            ej={'nombre': ' + '.join(names), 'subtitle':'\n'.join(notes), 'tipo':'circuito' if len(names)>1 else 'normal', 'planByMicro':{},'source':{'file':path.name,'sheet':ws.title,'row':start}}
            if len(names)>1: ej['circuitoLineas']=names
            for mn,c in enumerate(cols,1):
                cells=[ws.cell(r,cc) for r in range(start,kg) for cc in [c,c+1] if text(ws.cell(r,cc))]
                raw=[text(x) for x in cells]
                seriesText=text(ws.cell(start,c))
                match=re.match(r'^(\d+)(?:\s*o\s*(\d+))?\s*x\s*(.*)',seriesText,re.I)
                # Numeric series is storage capacity. Literal label preserves ranges and '?'.
                series=int(match[2] or match[1]) if match else 0
                slots=max([int(m[1]) for r in range(kg+1,end) if (m:=re.match(r'(\d+)ª Serie',text(ws.cell(r,1)),re.I))] or [0])
                reps=[]
                if match and match[3]: reps=[match[3]]
                for r in range(start,kg):
                    v=text(ws.cell(r,c+1))
                    if v and '%' not in v: reps.append(v)
                rest=[];restCells={}
                for r in range(kg+1,end):
                    v=text(ws.cell(r,c))
                    label=text(ws.cell(r,1))
                    if v and (re.search(r"\d.*(?:['\"′]|min)|alternar",v,re.I) or ('pausa' in label.lower() and v not in ['x','X'])):
                        rest.append(v);restCells[ws.cell(r,c).coordinate]=v
                percentages=[v for v in raw if '%' in v]
                ej['planByMicro'][str(mn)]={'series':series,'recordSlots':series or slots if raw else 0,'seriesLiteral':seriesText,'repsTarget':reps,'pausa':' / '.join(rest),'porcentaje':' / '.join(percentages),'label':' / '.join(raw),'indicaciones':en['indicaciones'][str(mn)],'sourceCells':{x.coordinate:text(x) for x in cells},'capacityUnspecified':bool(raw) and match is None}
                ej['planByMicro'][str(mn)]['sourceRestCells']=restCells
            en['ejercicios'].append(ej)
        data['entrenos'].append(en)
    # Maximales has two sessions, five attempts per exercise, not four repeated microcycles.
    if 'Maximales' in wb:
        ws=wb['Maximales']
        for session,(lo,hi) in enumerate([(1,35),(35,ws.max_row+1)],1):
            en={'letra':'MX'+str(session),'nombre':text(ws.cell(lo,1)),'numMicro':1,'ejercicios':[]}
            for r in range(lo,hi):
                for c in [1,7]:
                    name=text(ws.cell(r,c))
                    if not name or not re.match(r'1º intento',text(ws.cell(r+1,c))): continue
                    en['ejercicios'].append({'nombre':name,'subtitle':'\n'.join(text(ws.cell(rr,cc)) for rr in range(lo,hi) for cc in [1,7] if len(text(ws.cell(rr,cc)))>150),'tipo':'normal','source':{'file':path.name,'sheet':ws.title,'row':r},'planByMicro':{'1':{'series':5,'repsTarget':[],'pausa':'','label':'5 intentos · maximal válido','porcentaje':''}}})
            for r in range(lo,hi):
                m=re.match(r'(\d+) x',text(ws.cell(r,9)))
                if m and text(ws.cell(r,7)):
                    en['ejercicios'].append({'nombre':text(ws.cell(r,7)),'subtitle':'','tipo':'normal','source':{'file':path.name,'sheet':ws.title,'row':r},'planByMicro':{'1':{'series':int(m[1]),'repsTarget':[text(ws.cell(r,11))],'pausa':'','label':text(ws.cell(r,9))+' '+text(ws.cell(r,11)),'porcentaje':''}}})
            data['entrenos'].append(en)
    return data

def main():
    parser=argparse.ArgumentParser(); parser.add_argument('directory'); parser.add_argument('output'); args=parser.parse_args()
    data={'H':{},'M':{}}
    files=sorted(pathlib.Path(args.directory).glob('*Mesociclo*.xlsx'))
    assert len(files)==16, f'Expected 16 sources, got {len(files)}'
    for p in files:
        d=extract(p); cat=CATS[int(p.name[0])-1]; sex='M' if '(GIRL)' in p.name else 'H'; data[sex][cat]=d
        print(sex,cat,d['numMicro'],[(e['letra'],len(e['ejercicios'])) for e in d['entrenos']])
    pathlib.Path(args.output).write_text('/* Generated from original XLSX cells by scripts/import_biio.py. */\nconst TOB_BIIO_SOURCE = '+json.dumps(data,ensure_ascii=False,indent=2)+';\n',encoding='utf-8')

if __name__=='__main__': main()
